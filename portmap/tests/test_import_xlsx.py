import json
import sys
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.styles import PatternFill

APP = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP))
import import_xlsx as imp
import server as srv


def fill(cell, rgb):
    cell.fill = PatternFill("solid", fgColor="FF" + rgb)


def make_workbook(path):
    """A small copy of the real layout: switch blocks, the CRÍTICOS table and a
    rack grid on one sheet, desks with two side-by-side tables on another."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relação Rack - Switch nova"
    ws["B2"] = "SWITCH 1 (4 ports)"
    for i, (a, b) in enumerate([(1, 2), (3, 4)]):
        ws.cell(2 + i, 3, float(a))
        ws.cell(2 + i, 4, float(b))
    fill(ws["C2"], "000000")
    fill(ws["D2"], "000000")
    fill(ws["C3"], "00FF00")
    ws["B5"] = "SWITCH 2 (2 ports)"
    ws["C5"], ws["D5"] = 1.0, 2.0
    fill(ws["D5"], "FFD966")
    fill(ws["C5"], "123456")

    ws["G2"], ws["I2"], ws["K2"] = "NOME", "RACK", "SWITCH (nº, porta)"
    rows = [
        ("Sala de vidro WIFI", "R1 PP1 T02 ", "(1, 3)"),
        ("Link Gateway", "x", "(1, 1) (2,1)"),
        ("Impressora", "R1 PP2 T01", "(2, 2)"),
        ("Fantasma", "R1 PP1 T01", "não identificado"),
    ]
    for i, (name, rack, sw) in enumerate(rows):
        ws.cell(3 + i, 7, name)
        ws.cell(3 + i, 9, rack)
        ws.cell(3 + i, 11, sw)

    ws["C10"], ws["D10"] = "PP1", "PP2"
    ws["B11"] = "RACK 1 "
    for i in range(3):
        ws.cell(11 + i, 3, f"T{i + 1:02d}")
        ws.cell(11 + i, 4, f"T{i + 1:02d}")
    fill(ws["C12"], "4A86E8")
    fill(ws["D11"], "FF0000")

    ref = wb.create_sheet("referência pré-mudança ")
    for col, v in enumerate(["Setor", "Referência", "Baia", "Letra", "Endpoint 1", "Endpoint 2", "Status Infra cabeada:"], 1):
        ref.cell(3, col, v)
    for col, v in enumerate(["Local", "Endpoint 1", "Endpoint 2", "Status Infra cabeada:"], 9):
        ref.cell(3, col, v)
    data = [
        ["Receitas", "Mesa 1", 1.0, "A", "R1 PP1 T02", "r1pp1t3", "Íntegra"],
        ["Receitas", "Mesa 1", 1.0, "B", "x", "caso especial", "parcialmente danificada "],
    ]
    for r, row in enumerate(data, 4):
        for c, v in enumerate(row, 1):
            ref.cell(r, c, v)
    ref["I4"], ref["J4"], ref["K4"], ref["L4"] = "Sala 1", "R1 PP2 T01", "x", "íntegra"
    ref["I5"], ref["J5"], ref["K5"], ref["L5"] = "Setor", "Referência", "Endpoint", "Status infra cabeada:"
    ref["I6"], ref["J6"], ref["K6"], ref["L6"] = "Operações", "Impressora", "R1 PP2 T01", "Íntegra"
    wb.save(path)


@pytest.fixture
def imported(tmp_path):
    xlsx = tmp_path / "infra.xlsx"
    make_workbook(xlsx)
    out = tmp_path / "network.json"
    assert imp.main([str(xlsx), "-o", str(out), "--site", "Test"]) == 0
    return json.loads(out.read_text(encoding="utf-8")), tmp_path, xlsx, out


def test_switch_ports_and_states(imported):
    doc, *_ = imported
    assert srv.validate(doc) == []
    sw1, sw2 = doc["devices"]
    assert (sw1["id"], sw1["ports"], sw2["ports"]) == ("sw1", 4, 2)
    assert sw1["port_data"]["1"]["state"] == "gateway"
    assert sw1["port_data"]["3"]["state"] == "link_activity"
    assert "4" not in sw1["port_data"]
    assert sw2["port_data"]["2"]["state"] == "link_only"
    assert sw2["port_data"]["1"]["state"] == "unidentified"


def test_critical_table(imported):
    doc, *_ = imported
    sw1, sw2 = doc["devices"]
    p = sw1["port_data"]["3"]
    assert (p["name"], p["connected"], p["patch"], p["critical"]) == ("Sala de vidro WIFI", "Access point", "R1 PP1 T02", True)
    assert sw1["port_data"]["1"]["name"] == sw2["port_data"]["1"]["name"] == "Link Gateway"
    assert "patch" not in sw1["port_data"]["1"]
    assert sw2["port_data"]["2"]["connected"] == "Printer"


def test_override_note(imported):
    doc, *_ = imported
    p = doc["devices"][0]["port_data"]["2"]
    assert p["state"] == "gateway" and p["name"].startswith("Uplink") and "not recorded" in p["note"]


def test_rack_grid(imported):
    doc, *_ = imported
    rack = doc["racks"][0]
    assert (rack["id"], rack["panels"], rack["ports_per_panel"]) == ("R1", 2, 3)
    assert rack["cells"] == {"PP1 T02": {"state": "full_chain"}, "PP2 T01": {"state": "damaged"}}


def test_desks(imported):
    doc, *_ = imported
    eps = doc["endpoints"]
    assert {"patch": "R1 PP1 T02", "setor": "Receitas", "referencia": "Mesa 1", "baia": "1", "letra": "A",
            "cabling": "Íntegra", "jack": "Endpoint 1"} in eps
    assert any(e["patch"] == "R1 PP1 T03" and e["jack"] == "Endpoint 2" for e in eps)
    assert {"patch": "R1 PP2 T01", "referencia": "Sala 1", "cabling": "Íntegra", "jack": "Endpoint 1"} in eps
    assert {"patch": "R1 PP2 T01", "setor": "Operações", "referencia": "Impressora", "cabling": "Íntegra",
            "jack": "Endpoint"} in eps
    assert len(eps) == 4


def test_warnings_and_overwrite(imported, capsys):
    _, tmp_path, xlsx, out = imported
    with pytest.raises(SystemExit):
        imp.main([str(xlsx), "-o", str(out)])
    capsys.readouterr()
    assert imp.main([str(xlsx), "-o", str(out), "--force"]) == 0
    err = capsys.readouterr().err
    assert "fill 123456" in err
    assert "'caso especial'" in err
    assert "'Fantasma' has no switch port" in err
    assert json.loads(out.read_text(encoding="utf-8"))["revision"] == 2
    assert len(list((tmp_path / "backups").glob("*.json"))) == 1
