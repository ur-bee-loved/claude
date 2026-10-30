#!/usr/bin/env python3
"""Build data/network.json from the infrastructure spreadsheet.

Reads, by default:
  "Relação Rack - Switch nova"   switch ports (fill color = state), the
                                 CRÍTICOS table (name, rack position, switch
                                 port) and the rack grids (fill color = state)
  "referência pré-mudança"       desks (Setor, Referência, Baia, Letra) and the
                                 patch-panel positions their jacks use

Tables are found by their header text, not by fixed cell addresses. Requires
openpyxl (pip install openpyxl); the server itself does not.
"""

import argparse
import datetime as dt
import re
import sys
from pathlib import Path

from server import DATA_FILE, PROG, SCHEMA_VERSION, Store, validate

STATES = [
    {"id": "full_chain", "label": "Cadeia completa funcional", "color": "#3d8bff", "pattern": "solid"},
    {"id": "link_activity", "label": "Link E atividade", "color": "#2bff88", "pattern": "solid"},
    {"id": "link_only", "label": "Apenas link", "color": "#ffb000", "pattern": "solid"},
    {"id": "available", "label": "Disponível", "color": "#8a96ab", "pattern": "dashed"},
    {"id": "damaged", "label": "Danificado", "color": "#ff3355", "pattern": "solid"},
    {"id": "vital", "label": "Infra vital", "color": "#ff2bd6", "pattern": "solid"},
    {"id": "gateway", "label": "Link Gateway", "color": "#e8eef8", "pattern": "solid"},
    {"id": "unidentified", "label": "Não identificado", "color": "#31d7e8", "pattern": "hatched"},
]

# Spreadsheet fill (RRGGBB) -> state id, from the legend in the rack sheet.
FILL_TO_STATE = {
    None: "available",
    "FFFFFF": "available",
    "4A86E8": "full_chain",
    "00FF00": "link_activity",
    "F1C232": "link_only",
    "FFD966": "link_only",  # pale yellow on Switch 2, confirmed as "Apenas link"
    "CC0000": "damaged",
    "FF0000": "damaged",  # bright red rack cells (PESA, MUNIFE), confirmed as "Danificado"
    "FF00FF": "vital",
    "000000": "gateway",
    "00FFFF": "unidentified",
}

# First matching keyword in the reference name sets "Connected device".
KIND_RULES = [
    ("wifi", "Access point"),
    ("impressora", "Printer"),
    ("switch", "Switch"),
    ("gateway", "Gateway"),
    ("sala", "Room outlet"),
]

# Facts confirmed outside the spreadsheet, applied after reading it.
OVERRIDES = {
    ("sw1", 2): {
        "name": "Uplink Switch 1 ↔ Switch 2",
        "connected": "Switch",
        "note": "Marked black (Link Gateway) in the sheet, but the CRÍTICOS list only names (1,1) and (2,1) "
                "as gateway links. Confirmed as the uplink between the two switches; the Switch 2 port at the "
                "other end is not recorded.",
    },
}

SWITCH_RE = re.compile(r"SWITCH\s*(\d+)\s*\(\s*(\d+)\s*ports?\s*\)", re.IGNORECASE)
RACK_RE = re.compile(r"^RACK\s*(\d+)$", re.IGNORECASE)
PATCH_RE = re.compile(r"R\s*(\d+)\s*PP\s*(\d+)\s*T\s*(\d+)", re.IGNORECASE)
SWITCH_REF_RE = re.compile(r"\(\s*(\d+)\s*,\s*(\d+)\s*\)")
HEADER_ROLES = {"setor": "setor", "referência": "referencia", "referencia": "referencia", "local": "referencia",
                "baia": "baia", "letra": "letra"}


def die(msg, code=1):
    print(f"{PROG} import: error: {msg}", file=sys.stderr)
    sys.exit(code)


def text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def fill_hex(cell):
    f = cell.fill
    if not f or not f.fill_type or f.fgColor is None or f.fgColor.type != "rgb":
        return None
    return str(f.fgColor.rgb)[-6:].upper()


def parse_patch(v):
    m = PATCH_RE.fullmatch(text(v))
    return f"R{int(m[1])} PP{int(m[2])} T{int(m[3]):02d}" if m else None


def kind_of(name):
    low = name.lower()
    return next((kind for key, kind in KIND_RULES if key in low), "")


class Importer:
    def __init__(self):
        self.warnings = []

    def warn(self, msg):
        self.warnings.append(msg)

    def state(self, cell, where):
        hx = fill_hex(cell)
        if hx in FILL_TO_STATE:
            return FILL_TO_STATE[hx]
        self.warn(f"{where}: fill {hx} is not in the legend, marked Não identificado")
        return "unidentified"

    def switches(self, ws):
        devices = []
        for row in ws.iter_rows():
            for cell in row:
                m = SWITCH_RE.search(text(cell.value))
                if not m:
                    continue
                num, count = int(m[1]), int(m[2])
                dev = {"id": f"sw{num}", "type": "switch", "name": f"Switch {num}", "ports": count, "port_data": {}}
                r = cell.row
                while True:
                    found = False
                    for c in (cell.column + 1, cell.column + 2):
                        pc = ws.cell(r, c)
                        if isinstance(pc.value, (int, float)) and 1 <= pc.value <= count:
                            found = True
                            st = self.state(pc, f"{ws.title}!{pc.coordinate} (switch {num} port {int(pc.value)})")
                            if st != "available":
                                dev["port_data"][str(int(pc.value))] = {"state": st}
                    if not found:
                        break
                    r += 1
                devices.append(dev)
        return devices

    def critical(self, ws, devices):
        by_id = {d["id"]: d for d in devices}
        header = next((c for row in ws.iter_rows() for c in row if text(c.value).upper() == "NOME"), None)
        if header is None:
            self.warn(f"{ws.title}: no NOME column, critical devices skipped")
            return
        row_cells = {text(ws.cell(header.row, c).value).upper(): c for c in range(header.column, header.column + 12)}
        rack_col = row_cells.get("RACK")
        sw_col = next((c for k, c in row_cells.items() if k.startswith("SWITCH")), None)
        r = header.row + 1
        while (name := text(ws.cell(r, header.column).value)):
            patch = parse_patch(ws.cell(r, rack_col).value) if rack_col else None
            refs = SWITCH_REF_RE.findall(text(ws.cell(r, sw_col).value)) if sw_col else []
            if not refs:
                self.warn(f"{ws.title}!row {r}: {name!r} has no switch port, skipped")
            for sw, port in refs:
                dev = by_id.get(f"sw{sw}")
                if not dev or not 1 <= int(port) <= dev["ports"]:
                    self.warn(f"{ws.title}!row {r}: {name!r} points at missing port ({sw}, {port})")
                    continue
                p = dev["port_data"].setdefault(port, {"state": "available"})
                if p.get("name"):
                    self.warn(f"switch {sw} port {port}: {p['name']!r} replaced by {name!r}")
                p.update({"name": name, "critical": True})
                if kind := kind_of(name):
                    p["connected"] = kind
                if patch:
                    p["patch"] = patch
            r += 1

    def racks(self, ws):
        racks = []
        for row in ws.iter_rows():
            for cell in row:
                m = RACK_RE.match(text(cell.value))
                if not m:
                    continue
                rid = f"R{int(m[1])}"
                rack = {"id": rid, "name": f"Rack {int(m[1])}", "panels": 0, "ports_per_panel": 0, "cells": {}}
                panels = {}
                c = cell.column + 1
                while (pm := re.fullmatch(r"PP(\d+)", text(ws.cell(cell.row - 1, c).value), re.IGNORECASE)):
                    panels[c] = int(pm[1])
                    c += 1
                r = cell.row
                while (tm := re.fullmatch(r"T(\d+)", text(ws.cell(r, cell.column + 1).value), re.IGNORECASE)):
                    port = int(tm[1])
                    rack["ports_per_panel"] = max(rack["ports_per_panel"], port)
                    for col, pp in panels.items():
                        pc = ws.cell(r, col)
                        st = self.state(pc, f"{ws.title}!{pc.coordinate} ({rid} PP{pp} T{port:02d})")
                        if st != "available":
                            rack["cells"][f"PP{pp} T{port:02d}"] = {"state": st}
                    r += 1
                rack["panels"] = max(panels.values(), default=0)
                if rack["panels"] and rack["ports_per_panel"]:
                    racks.append(rack)
                else:
                    self.warn(f"{ws.title}!{cell.coordinate}: {rid} has no PP/T grid, skipped")
        return racks

    def endpoints(self, ws):
        """Rows of every table whose header has Endpoint columns. Header cells set
        the role of their column until another header row changes it."""
        roles = {}
        out = []
        for row in ws.iter_rows():
            header_cols = set()
            for cell in row:
                t = text(cell.value).lower().rstrip(":")
                role = HEADER_ROLES.get(t) or ("endpoint" if t.startswith("endpoint") else None) \
                    or ("cabling" if t.startswith("status") else None)
                if role:
                    roles[cell.column] = (role, text(cell.value))
                    header_cols.add(cell.column)
            for run in self._runs(sorted(roles)):
                if header_cols & set(run):
                    continue
                values = {col: ws.cell(row[0].row, col).value for col in run}
                base = {}
                for col in run:
                    role, _ = roles[col]
                    if role in ("setor", "referencia", "baia", "letra", "cabling") and text(values[col]):
                        v = text(values[col])
                        base[role] = v[:1].upper() + v[1:] if role == "cabling" else v
                for col in run:
                    role, label = roles[col]
                    if role != "endpoint":
                        continue
                    raw = text(values[col])
                    patch = parse_patch(raw)
                    if patch:
                        out.append({"patch": patch, **base, "jack": label})
                    elif raw and raw.lower() != "x":
                        self.warn(f"{ws.title}!row {row[0].row}: endpoint {raw!r} is not a patch reference, skipped")
        return out

    @staticmethod
    def _runs(cols):
        runs, cur = [], []
        for c in cols:
            if cur and c != cur[-1] + 1:
                runs.append(cur)
                cur = []
            cur.append(c)
        if cur:
            runs.append(cur)
        return runs


def find_sheet(wb, name):
    want = name.strip().casefold()
    for ws in wb.worksheets:
        if ws.title.strip().casefold() == want:
            return ws
    die(f"sheet {name!r} not found; sheets: {', '.join(repr(w.title) for w in wb.worksheets)}")


def build(wb, switch_sheet, desk_sheet, site, source_name):
    imp = Importer()
    ws = find_sheet(wb, switch_sheet)
    devices = imp.switches(ws)
    if not devices:
        die(f"no 'SWITCH n (N ports)' block in {ws.title!r}")
    imp.critical(ws, devices)
    for (dev_id, port), values in OVERRIDES.items():
        dev = next((d for d in devices if d["id"] == dev_id), None)
        if dev and port <= dev["ports"]:
            dev["port_data"].setdefault(str(port), {"state": "available"}).update(values)
    racks = imp.racks(ws)
    endpoints = imp.endpoints(find_sheet(wb, desk_sheet)) if desk_sheet else []
    doc = {
        "version": SCHEMA_VERSION,
        "revision": 0,
        "site": site,
        "source": {"file": source_name, "imported_utc": dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
                   "sheets": [s for s in (switch_sheet, desk_sheet) if s]},
        "states": STATES,
        "devices": devices,
        "racks": racks,
        "endpoints": endpoints,
    }
    return doc, imp.warnings


def main(argv=None):
    p = argparse.ArgumentParser(prog=f"{PROG} import", description="Build the portmap data file from the spreadsheet.")
    p.add_argument("--h", action="help", help=argparse.SUPPRESS)
    p.add_argument("xlsx", type=Path)
    p.add_argument("-o", "--output", type=Path, default=DATA_FILE, help="default data/network.json")
    p.add_argument("--switch-sheet", default="Relação Rack - Switch nova")
    p.add_argument("--desk-sheet", default="referência pré-mudança", help="'' to skip desks")
    p.add_argument("--site", default="Infra", help="site name shown in the header")
    p.add_argument("-f", "--force", action="store_true", help="replace an existing data file (a backup is kept)")
    args = p.parse_args(argv)

    try:
        import openpyxl
    except ImportError:
        die("openpyxl is required: pip install openpyxl")
    if not args.xlsx.is_file():
        die(f"{args.xlsx}: no such file")
    if args.output.exists() and not args.force:
        die(f"{args.output} exists; use --force to replace it (the old file goes to {args.output.parent / 'backups'})")

    wb = openpyxl.load_workbook(args.xlsx)
    doc, warnings = build(wb, args.switch_sheet, args.desk_sheet, args.site, args.xlsx.name)
    problems = validate(doc)
    if problems:
        die("generated data is invalid: " + "; ".join(problems))

    store = Store(args.output, args.output)
    if args.output.exists():
        doc["revision"] = store.load()[0].get("revision", 0)
    store.save(doc)

    for w in warnings:
        print(f"warning: {w}", file=sys.stderr)
    named = sum(1 for d in doc["devices"] for x in d["port_data"].values() if x.get("name"))
    print(f"{args.output}: {len(doc['devices'])} devices ({named} named ports), {len(doc['racks'])} racks "
          f"({sum(len(r['cells']) for r in doc['racks'])} marked positions), {len(doc['endpoints'])} desk jacks, "
          f"{len(warnings)} warnings", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
